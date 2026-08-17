"""Reading the Matrixify export and writing the enriched result back out."""

import html
import re

import pandas as pd
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

from .config import METAFIELD_ABOUT_THIS_ITEM, METAFIELD_SPECIFICATIONS, OUTPUT_COLUMNS

ASIN_RE = re.compile(r"^[A-Z0-9]{10}$")


def _pick_sheet(path, sheet):
    if sheet:
        return sheet
    xl = pd.ExcelFile(path, engine="openpyxl")
    for name in xl.sheet_names:
        if name.strip().lower() == "products":
            return name
    return xl.sheet_names[0]


def detect_asin_column(columns, override=None):
    if override:
        for col in columns:
            if col.strip().lower() == override.strip().lower():
                return col
        raise ValueError(
            f"Column '{override}' not found. Available columns: {list(columns)}"
        )

    candidates = [c for c in columns if "asin" in str(c).strip().lower()]
    if not candidates:
        raise ValueError(
            "Could not find a column with 'ASIN' in its name. "
            f"Available columns: {list(columns)}. "
            "Pass --asin-column to specify it explicitly."
        )
    exact = [c for c in candidates if str(c).strip().lower() == "asin"]
    return exact[0] if exact else candidates[0]


def read_input(path, sheet=None, asin_column=None):
    sheet_name = _pick_sheet(path, sheet)
    df = pd.read_excel(path, sheet_name=sheet_name, engine="openpyxl", dtype=str)
    df = df.dropna(how="all")
    col = detect_asin_column(df.columns, asin_column)

    df[col] = df[col].astype(str).str.strip().str.upper()
    invalid = df[~df[col].str.match(ASIN_RE, na=False)]
    if len(invalid):
        bad_preview = invalid[col].head(5).tolist()
        print(
            f"Warning: {len(invalid)} row(s) have a value in '{col}' that doesn't "
            f"look like a valid ASIN (e.g. {bad_preview}). They will be skipped."
        )
    df = df[df[col].str.match(ASIN_RE, na=False)].reset_index(drop=True)
    return df, col, sheet_name


def find_body_html_column(columns):
    for col in columns:
        name = str(col).strip().lower()
        if "body" in name and "html" in name:
            return col
    return None


def find_title_column(columns):
    for col in columns:
        if str(col).strip().lower() == "title":
            return col
    return None


def find_condition_column(columns):
    for col in columns:
        if str(col).strip().lower() == "condition":
            return col
    return None


def _fill_if_blank(existing_value, new_value):
    existing = str(existing_value or "").strip()
    if existing.lower() in ("nan", "none"):
        existing = ""
    return existing or new_value


def _build_condition_badge(condition_value):
    condition = str(condition_value or "").strip()
    if condition.lower() in ("", "nan", "none"):
        return ""
    return (
        '<p><span style="display:inline-block;padding:4px 14px;border-radius:999px;'
        "background:#eef0fe;color:#4f46e5;font-weight:600;font-size:13px;"
        'font-family:sans-serif;">'
        f"Condition: {html.escape(condition)}</span></p>"
    )


def build_output_rows(df, asin_col, results_by_asin, body_col, title_col, condition_col):
    rows = []
    for _, row in df.iterrows():
        asin = row[asin_col]
        result = results_by_asin.get(asin, {})
        out = dict(row)

        existing_body = str(row.get(body_col, "") or "").strip()
        if existing_body.lower() in ("nan", "none"):
            existing_body = ""
        new_body = result.get("body_html", "")
        if new_body and existing_body:
            merged_body = existing_body + "\n<hr/>\n" + new_body
        else:
            merged_body = new_body or existing_body

        if condition_col:
            badge = _build_condition_badge(row.get(condition_col))
            if badge:
                merged_body = f"{badge}\n{merged_body}" if merged_body else badge

        out.update(
            {
                "Image Src": result.get("image_src", ""),
                title_col: _fill_if_blank(row.get(title_col), result.get("title", "")),
                METAFIELD_ABOUT_THIS_ITEM: result.get("about_this_item", ""),
                METAFIELD_SPECIFICATIONS: result.get("specifications", ""),
                body_col: merged_body,
            }
        )
        rows.append(out)
    return rows


def write_output(df, asin_col, results_by_asin, output_path):
    body_col = find_body_html_column(df.columns) or "Body HTML"
    title_col = find_title_column(df.columns) or "Title"
    condition_col = find_condition_column(df.columns)
    rows = build_output_rows(df, asin_col, results_by_asin, body_col, title_col, condition_col)
    out_df = pd.DataFrame(rows)

    original_cols = list(df.columns)
    for extra_col in (body_col, title_col):
        if extra_col not in original_cols:
            original_cols.append(extra_col)
    ordered_cols = [c for c in original_cols if c not in OUTPUT_COLUMNS] + [
        c for c in OUTPUT_COLUMNS if c in out_df.columns
    ]
    out_df = out_df[ordered_cols]

    out_df.to_excel(output_path, index=False, sheet_name="Products", engine="openpyxl")
    _format_output(output_path, ordered_cols, body_col, title_col)
    return output_path


def _format_output(path, columns, body_col, title_col):
    from openpyxl import load_workbook

    wb = load_workbook(path)
    ws = wb.active
    ws.freeze_panes = "A2"

    wrap_cols = {"Image Src", METAFIELD_ABOUT_THIS_ITEM, METAFIELD_SPECIFICATIONS, body_col}
    widths = {
        "Image Src": 60,
        METAFIELD_ABOUT_THIS_ITEM: 60,
        METAFIELD_SPECIFICATIONS: 60,
        title_col: 40,
        body_col: 60,
    }

    for idx, col_name in enumerate(columns, start=1):
        letter = get_column_letter(idx)
        ws.column_dimensions[letter].width = widths.get(col_name, 16)
        if col_name in wrap_cols:
            for cell in ws[letter][1:]:
                cell.alignment = Alignment(wrap_text=True, vertical="top")

    for cell in ws[1]:
        cell.font = Font(bold=True)

    wb.save(path)
